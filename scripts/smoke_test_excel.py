"""Offline smoke test for workbook generation (does not need PaddleOCR models)."""
from pathlib import Path
import tempfile
import sys

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook
from core.excel_writer import write_workbook


class FakeResult:
    @property
    def json(self):
        return {
            "res": {
                "overall_ocr_res": {
                    "rec_texts": ["金额", " 123.45\n", "公式 =1+1"],
                    "rec_scores": [0.99, 0.9723, 0.91],
                    "rec_boxes": [[1, 2, 30, 20], [32, 2, 70, 20], [1, 25, 60, 42]],
                },
                "table_res_list": [{
                    "pred_html": "<table><tr><th>字段</th><th>值</th></tr><tr><td rowspan='2'>总计</td><td>123</td></tr><tr><td>456</td></tr></table>",
                    "table_ocr_pred": {
                        "rec_texts": ["字段", "值", "总计", "123", "456"],
                        "rec_scores": [0.99, 0.99, 0.95, 0.98, 0.96],
                        "rec_boxes": [],
                    },
                }],
            }
        }


def main():
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "smoke.xlsx"
        stats = write_workbook(out, [{"filename": "sample.png", "result": FakeResult()}], remove_whitespace=True)
        wb = load_workbook(out, data_only=False)
        assert wb.sheetnames == ["Sheet1_表格结果", "Sheet2_文字置信度"], wb.sheetnames
        assert stats["tables"] == 1, stats
        assert wb["Sheet1_表格结果"]["A2"].value == "字段", wb["Sheet1_表格结果"]["A2"].value
        assert wb["Sheet2_文字置信度"].max_row >= 8
        # OCR text that looks like a formula must be kept as a literal string.
        formula_values = [cell for row in wb["Sheet2_文字置信度"].iter_rows() for cell in row if cell.value == "公式=1+1"]
        assert formula_values, "whitespace cleaning or literal text write failed"
        assert formula_values[0].data_type == "s", formula_values[0].data_type
        print("SMOKE TEST OK", stats, wb.sheetnames)


if __name__ == "__main__":
    main()
