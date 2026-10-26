from __future__ import annotations

import logging
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .ocr_engine import result_to_dict
from .preprocess import clean_text

logger = logging.getLogger(__name__)

_HEADER_FILL = PatternFill("solid", fgColor="E8EEF9")
_SECTION_FILL = PatternFill("solid", fgColor="DCE7F8")
_THIN_GREY = Side(style="thin", color="D5DCE7")
_CELL_BORDER = Border(left=_THIN_GREY, right=_THIN_GREY, top=_THIN_GREY, bottom=_THIN_GREY)


def _string(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def _seq(value: Any) -> list[Any]:
    if value is None:
        return []
    if isinstance(value, list):
        return value
    if isinstance(value, tuple):
        return list(value)
    if hasattr(value, "tolist"):
        try:
            casted = value.tolist()
            return casted if isinstance(casted, list) else [casted]
        except Exception:
            pass
    return [value]


def _safe_cell(cell: Any, value: Any) -> None:
    """Write OCR strings as literal text, never as Excel formulas."""
    if isinstance(value, str):
        cell.value = value
        if value.startswith(("=", "+", "-", "@")):
            cell.data_type = "s"
    else:
        cell.value = value
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def extract_html_tables(payload: dict[str, Any]) -> list[str]:
    tables: list[str] = []
    for table in _seq(payload.get("table_res_list")):
        if isinstance(table, dict):
            html = table.get("pred_html") or table.get("html")
            if isinstance(html, str) and "<table" in html.lower():
                tables.append(html)

    # Compatibility fallback for result variants where table HTML is embedded in layout blocks.
    if not tables:
        for block in _seq(payload.get("parsing_res_list")):
            if not isinstance(block, dict):
                continue
            content = block.get("block_content")
            if isinstance(content, str) and "<table" in content.lower():
                tables.append(content)
    return tables


def html_table_to_grid(html: str, remove_whitespace: bool = False) -> tuple[list[list[str]], list[tuple[int, int, int, int]]]:
    """Read a predicted HTML table while preserving rowspan/colspan as Excel merges."""
    soup = BeautifulSoup(html, "lxml")
    table = soup.find("table")
    if table is None:
        return [], []
    rows = table.find_all("tr")
    if not rows:
        return [], []

    cells: dict[tuple[int, int], str] = {}
    reserved: set[tuple[int, int]] = set()
    merges: list[tuple[int, int, int, int]] = []  # zero-based r1,c1,r2,c2
    max_col = 0

    for row_index, row in enumerate(rows):
        col_index = 0
        row_cells = row.find_all(["td", "th"], recursive=False)
        # Some HTML serializers wrap table cells inside tbody/tr normally; recursive=False
        # on tr is sufficient, but retain a fallback for malformed model-generated HTML.
        if not row_cells:
            row_cells = row.find_all(["td", "th"])
        for cell in row_cells:
            while (row_index, col_index) in reserved:
                col_index += 1
            try:
                rowspan = max(1, int(cell.get("rowspan", 1)))
            except (TypeError, ValueError):
                rowspan = 1
            try:
                colspan = max(1, int(cell.get("colspan", 1)))
            except (TypeError, ValueError):
                colspan = 1

            value = " ".join(cell.stripped_strings)
            value = clean_text(value, remove_whitespace)
            cells[(row_index, col_index)] = value

            if rowspan > 1 or colspan > 1:
                merges.append((row_index, col_index, row_index + rowspan - 1, col_index + colspan - 1))
            for rr in range(row_index, row_index + rowspan):
                for cc in range(col_index, col_index + colspan):
                    if (rr, cc) != (row_index, col_index):
                        reserved.add((rr, cc))
            col_index += colspan
            max_col = max(max_col, col_index)

    max_row = len(rows)
    if reserved:
        max_row = max(max_row, max(r for r, _ in reserved) + 1)
        max_col = max(max_col, max(c for _, c in reserved) + 1)
    grid = [[cells.get((r, c), "") for c in range(max_col)] for r in range(max_row)]
    return grid, merges


def extract_text_records(filename: str, result: Any, remove_whitespace: bool = False) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    """Return normalized result payload plus all OCR text/score records."""
    if isinstance(result, dict) and result.get("_error"):
        return {}, [{
            "filename": filename,
            "source": "处理异常",
            "index": 1,
            "text": str(result["_error"]),
            "score": None,
            "box": "",
        }]

    payload = result_to_dict(result)
    records: list[dict[str, Any]] = []

    def append_from_ocr(ocr: Any, source: str) -> None:
        if not isinstance(ocr, dict):
            return
        texts = _seq(ocr.get("rec_texts"))
        scores = _seq(ocr.get("rec_scores"))
        boxes = _seq(ocr.get("rec_boxes"))
        if not boxes:
            boxes = _seq(ocr.get("dt_polys"))
        for idx, text in enumerate(texts):
            score = scores[idx] if idx < len(scores) else None
            box = boxes[idx] if idx < len(boxes) else ""
            try:
                score = float(score) if score is not None else None
            except (TypeError, ValueError):
                score = None
            records.append({
                "filename": filename,
                "source": source,
                "index": idx + 1,
                "text": clean_text(_string(text), remove_whitespace),
                "score": score,
                "box": _string(box),
            })

    append_from_ocr(payload.get("overall_ocr_res"), "整页OCR")
    for table_index, table in enumerate(_seq(payload.get("table_res_list")), start=1):
        if isinstance(table, dict):
            append_from_ocr(table.get("table_ocr_pred"), f"表格{table_index}单元格OCR")

    # Fallback: if structured OCR lists are not exposed by a particular build, preserve
    # readable layout content even when that content does not come with a confidence score.
    if not records:
        for block_index, block in enumerate(_seq(payload.get("parsing_res_list")), start=1):
            if not isinstance(block, dict):
                continue
            content = block.get("block_content")
            if isinstance(content, str) and content.strip() and "<table" not in content.lower():
                records.append({
                    "filename": filename,
                    "source": "版面文本(模型未返回分数)",
                    "index": block_index,
                    "text": clean_text(content, remove_whitespace),
                    "score": None,
                    "box": _string(block.get("block_bbox", "")),
                })

    return payload, records


def write_workbook(
    output_path: str | Path,
    result_items: list[dict[str, Any]],
    remove_whitespace: bool = False,
) -> dict[str, int]:
    """Create a two-sheet workbook: recognized tables and OCR text with confidence."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws_tables = wb.active
    ws_tables.title = "Sheet1_表格结果"
    ws_text = wb.create_sheet("Sheet2_文字置信度")

    ws_text_headers = ["图片文件名", "来源", "序号", "识别文本", "置信度", "文本框坐标"]
    ws_text.append(ws_text_headers)
    for cell in ws_text[1]:
        cell.font = Font(bold=True, color="22304A")
        cell.fill = _HEADER_FILL
        cell.border = _CELL_BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws_text.freeze_panes = "A2"

    row = 1
    total_tables = 0
    total_texts = 0
    for item in result_items:
        filename = str(item.get("filename", "unknown"))
        error = item.get("error")
        result = item.get("result")
        if error:
            payload: dict[str, Any] = {}
            records = [{
                "filename": filename,
                "source": "处理异常",
                "index": 1,
                "text": str(error),
                "score": None,
                "box": "",
            }]
        else:
            try:
                payload, records = extract_text_records(filename, result, remove_whitespace)
            except Exception as exc:
                logger.exception("解析识别结果失败 | 文件=%s", filename)
                payload, records = {}, [{
                    "filename": filename,
                    "source": "结果解析异常",
                    "index": 1,
                    "text": f"{type(exc).__name__}: {exc}",
                    "score": None,
                    "box": "",
                }]

        for rec in records:
            ws_text.append([
                rec["filename"], rec["source"], rec["index"], rec["text"], rec["score"], rec["box"],
            ])
            excel_row = ws_text.max_row
            for cell in ws_text[excel_row]:
                cell.border = _CELL_BORDER
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            if rec["score"] is not None:
                ws_text.cell(excel_row, 5).number_format = "0.0000"
            total_texts += 1

        table_htmls = extract_html_tables(payload)
        if table_htmls:
            for table_num, html in enumerate(table_htmls, start=1):
                grid, merges = html_table_to_grid(html, remove_whitespace)
                if not grid:
                    continue
                section_row = row
                row = section_row + 1
                max_table_cols = max(1, max((len(x) for x in grid), default=1))
                _safe_cell(ws_tables.cell(section_row, 1), f"图片：{filename} | 表格 {table_num}")
                ws_tables.cell(section_row, 1).font = Font(bold=True, size=11, color="22304A")
                ws_tables.cell(section_row, 1).fill = _SECTION_FILL
                if max_table_cols > 1:
                    ws_tables.merge_cells(start_row=section_row, start_column=1, end_row=section_row, end_column=max_table_cols)
                for r_idx, values in enumerate(grid):
                    for c_idx, value in enumerate(values, start=1):
                        cell = ws_tables.cell(row + r_idx, c_idx)
                        _safe_cell(cell, value)
                        cell.border = _CELL_BORDER
                        if r_idx == 0:
                            cell.font = Font(bold=True, color="22304A")
                            cell.fill = _HEADER_FILL
                for r1, c1, r2, c2 in merges:
                    if r2 > r1 or c2 > c1:
                        try:
                            ws_tables.merge_cells(
                                start_row=row + r1, start_column=c1 + 1,
                                end_row=row + r2, end_column=c2 + 1,
                            )
                        except ValueError:
                            logger.debug("跳过重叠的HTML单元格合并范围：%s", (r1, c1, r2, c2))
                row += len(grid) + 2
                total_tables += 1
        else:
            # Keep useful content on Sheet1 if the image is a text page or the detector missed the table.
            text_only = [r for r in records if r["source"] == "整页OCR" or "单元格OCR" in r["source"]]
            if text_only:
                _safe_cell(ws_tables.cell(row, 1), f"图片：{filename} | 未检测到表格，以下为文本识别结果")
                ws_tables.cell(row, 1).font = Font(bold=True, color="22304A")
                ws_tables.cell(row, 1).fill = _SECTION_FILL
                row += 1
                for col, title in enumerate(["识别文本", "可信度"], start=1):
                    cell = ws_tables.cell(row, col)
                    _safe_cell(cell, title)
                    cell.font = Font(bold=True)
                    cell.fill = _HEADER_FILL
                    cell.border = _CELL_BORDER
                row += 1
                for rec in text_only:
                    _safe_cell(ws_tables.cell(row, 1), rec["text"])
                    score_cell = ws_tables.cell(row, 2)
                    _safe_cell(score_cell, rec["score"])
                    score_cell.number_format = "0.0000"
                    ws_tables.cell(row, 1).border = _CELL_BORDER
                    score_cell.border = _CELL_BORDER
                    row += 1
                row += 2
            elif error or item.get("error"):
                _safe_cell(ws_tables.cell(row, 1), f"图片：{filename} | 处理失败：{error}")
                ws_tables.cell(row, 1).font = Font(color="9C2C2C")
                row += 2

    if row == 1:
        ws_tables.cell(1, 1, "未获得表格结果。请检查图片质量、离线模型包与日志。")

    # Table sheet ergonomics and readable width limits.
    ws_tables.freeze_panes = "A1"
    for col_cells in ws_tables.columns:
        col_letter = get_column_letter(col_cells[0].column)
        values = [len(str(c.value)) for c in col_cells if c.value is not None]
        ws_tables.column_dimensions[col_letter].width = min(max(max(values, default=12) + 2, 12), 34)
    for row_cells in ws_tables.iter_rows():
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    text_widths = {"A": 28, "B": 22, "C": 9, "D": 52, "E": 12, "F": 34}
    for col, width in text_widths.items():
        ws_text.column_dimensions[col].width = width
    ws_text.auto_filter.ref = ws_text.dimensions
    ws_text.row_dimensions[1].height = 28

    wb.properties.creator = "Offline OCR Excel - PaddleOCR PP-StructureV3"
    wb.properties.title = f"图片转Excel识别结果 {datetime.now():%Y-%m-%d %H:%M:%S}"
    wb.save(output_path)
    logger.info(
        "Excel生成完成 | 路径=%s | 表格数=%d | 文字记录数=%d | 总行数Sheet2=%d",
        output_path, total_tables, total_texts, ws_text.max_row,
    )
    return {"tables": total_tables, "text_records": total_texts, "images": len(result_items)}
